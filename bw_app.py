from time import sleep
import iperf3
import openpyxl
import tkinter as tk
import threading
from tkinter import messagebox
from tkinter.filedialog import asksaveasfilename
from tkinter import ttk
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import (FigureCanvasTkAgg, NavigationToolbar2Tk)
from openpyxl.chart import Reference, BarChart


Server_List = {'Singapore': '89.187.160.1', 'Tokyo': '89.187.162.1', 'HongKong': '84.17.57.129'}
Interation_List = {'10': 10, '20': 20, '50': 50, '100': 100, '1000': 1000}


class BandwidthTest(tk.Tk):
    def __init__(self, server='89.187.160.1', port=5201, duration=1, iterations=10):
        super().__init__()
        self.upl = []
        self.dowl = []
        self.server = server
        self.port = port
        self.duration = duration
        self.iterations = iterations
        self.test_results = []
        self.window = None
        self.ServerChoosen = None
        self.InterationChoosen = None
        self.progress = None
        self.title("Bandwidth Test")
        self.geometry("800x600")
        self.create_widget()

    def create_widget(self):

        menubar = tk.Menu(self)

        option = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label='option', menu=option)
        option.add_command(label="BW Test", command=self.bandwidth_test)
        option.add_command(label="Export as Excel", command=self.export_bandwidth_test_to_excel)
        option.add_command(label="Configure Server", command=self.configure_setting)
        option.add_separator()
        option.add_command(label="Exit", command=self.destroy)

        self.config(menu=menubar)

    def run_iperf3_test(self):
        client = iperf3.Client()
        client.server_hostname = self.server
        client.port = self.port
        client.duration = self.duration

        result = client.run()

        if result.error:
            #messagebox.showerror(title="result error", message=f"{result.error}")
            return
        else:
            return self.test_results.append({
                'sent_Mbps': result.sent_Mbps,
                'received_Mbps': result.received_Mbps,
            })

    def run_multiple_tests(self):
        self.test_results.clear()
        self.upl.clear()
        self.dowl.clear()
        for i in range(self.iterations):
            self.progress['value'] = i * (100 / self.iterations)
            self.window.update_idletasks()
            threading.Thread(target=self.run_iperf3_test()).start()
            sleep(1)

        self.progress['value'] = 100
        self.window.update_idletasks()
        error_cnt = 0

        for i, result in enumerate(self.test_results):
            if 'error' in result:
                error_cnt = error_cnt + 1
            else:
                self.upl.append(result["sent_Mbps"])
                self.dowl.append(result["received_Mbps"])

        if error_cnt >= (self.iterations/2):
            self.window.message = messagebox.showerror(title="test state", message="test failed")
            return
        self.window.message = messagebox.showinfo(title="test state", message="successfully test")
        fig = Figure(figsize=(5, 5), dpi=80)
        plot1 = fig.add_subplot(111)

        plot1.plot(self.upl, label='Upload')
        plot1.plot(self.dowl, label='Download')
        plot1.legend()

        canvas = FigureCanvasTkAgg(fig, master=self.window)
        canvas.draw()
        canvas.get_tk_widget().pack()

        toolbar = NavigationToolbar2Tk(canvas, self.window)
        toolbar.update()
        canvas.get_tk_widget().pack()

        self.window.result_text = tk.Text(self.window, height=10, width=50)
        self.window.result_text.pack()

        average_upl, average_dowl = self.average_bandwidth()

        self.window.result_text.insert(tk.END, f"Upload: {average_upl} Mbps\n"
                                               f"Dowload: {average_dowl} Mbps\n"
                                               f"Server: {self.server}\n")


    def bandwidth_test(self):
        self.window = tk.Tk()
        self.window.title("Bandwidth Test Module")
        self.window.geometry('800x600')
        self.progress = ttk.Progressbar(self.window, orient="horizontal", length=100, mode='determinate')
        self.progress.pack(pady=10)

        start_button = tk.Button(self.window, text='Start', command=self.run_multiple_tests)
        start_button.pack(pady=10)

    def export_bandwidth_test_to_excel(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Bandwidth Test Result"
        sheet['A1'].value = "Upload (Mbps)"
        sheet['B1'].value = "Download (Mbps)"
        row = 2
        for data in self.upl:
            sheet[f'A{row}'].value = "{:.3f}".format(data)
            row += 1

        row = 2
        for data in self.dowl:
            sheet[f'B{row}'].value = "{:.3f}".format(data)
            row += 1

        average_upl, average_dowl = self.average_bandwidth()
        sheet['D5'].value = "Average Upload: "
        sheet['D6'].value = "Average Download: "

        sheet['E5'].value = "{:.3f}".format(average_upl)
        sheet['E6'].value = "{:.3f}".format(average_dowl)

        # Create upload chart
        upload_chart = BarChart()
        upload_chart.title = "Upload Chart"
        upload_chart.x_axis.title = "Times"
        upload_chart.y_axis.title = "Mbps"
        upload_values = Reference(sheet, min_col=1, max_col=1, min_row=2, max_row=len(self.upl) + 1)
        upload_chart.add_data(upload_values, titles_from_data=False)
        sheet.add_chart(upload_chart, "G2")

        # Create download chart
        download_chart = BarChart()
        download_chart.title = "Download Chart"
        download_chart.x_axis.title = "Times"
        download_chart.y_axis.title = "Mbps"
        download_values = Reference(sheet, min_col=2, max_col=2, min_row=2, max_row=len(self.dowl) + 1)
        download_chart.add_data(download_values, titles_from_data=False)
        sheet.add_chart(download_chart, "G20")

        files = [('Excel Files', '*.xlsx')]
        save_path = asksaveasfilename(filetypes=files)
        if save_path:
            wb.save(save_path)
            messagebox.showinfo(title="Export state", message="Export Completed")

    def average_bandwidth(self):
        if len(self.upl) == 0 or len(self.dowl) == 0:
            return
        average_upl = sum(self.upl) / len(self.upl)
        average_dowl = sum(self.dowl) / len(self.dowl)
        return average_upl, average_dowl

    def configure_setting(self):
        window = tk.Tk()
        window.geometry('640x480')

        tk.Label(window, text="Select the Server :",
                 font=("Times New Roman", 14)).grid(column=0,
                                                    row=15, padx=10, pady=25)
        var = tk.StringVar()
        var2 = tk.StringVar()
        self.ServerChoosen = ttk.Combobox(window, width=27, textvariable=var)
        self.ServerChoosen['values'] = ('Singapore', 'Tokyo', 'HongKong')

        self.ServerChoosen.grid(column=1, row=15)
        self.ServerChoosen.current(0)

        tk.Label(window, text="Select Interation :",
                 font=("Times New Roman", 14)).grid(column=0,
                                                    row=40, padx=10, pady=25)
        self.InterationChoosen = ttk.Combobox(window, width=27, textvariable=var2)
        self.InterationChoosen['values'] = ('10', '20', '50', '100', '1000')

        self.InterationChoosen.grid(column=1, row=40)
        self.InterationChoosen.current(0)

        save_button = tk.Button(window, text="Save", command=self.save_seletion)
        save_button.grid(column=1, row=45, padx=10, pady=10)

    def save_seletion(self):
        choosen_server = self.ServerChoosen.get()
        self.server = Server_List['{}'.format(choosen_server)]

        choosen_interation = self.InterationChoosen.get()
        self.iterations = Interation_List['{}'.format(choosen_interation)]


if __name__ == '__main__':
    app = BandwidthTest()
    app.mainloop()