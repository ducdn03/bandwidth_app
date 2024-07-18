import json
import subprocess
import time
from time import sleep
import tkinter as tk
from functools import partial
from tkinter import messagebox
from tkinter import ttk
from tkinter.filedialog import asksaveasfilename

import openpyxl
from matplotlib.backends.backend_tkagg import (FigureCanvasTkAgg, NavigationToolbar2Tk)
from matplotlib.figure import Figure
from openpyxl.chart import Reference, LineChart

Server_List = {'Singapore': '89.187.160.1', 'Tokyo': 'speedtest.tyo11.jp.leaseweb.net', 'HongKong': '84.17.57.129'}
Interation_List = {'1': 1, '10': 10, '20': 20, '50': 50, '100': 100, '1000': 1000}


class BandwidthTest(tk.Tk):
    def __init__(self, server='89.187.160.1', port=5201, duration=10, iterations=1):
        super().__init__()
        self.upl = []
        self.dowl = []
        self.server = server
        self.port = port
        self.duration = duration
        self.iterations = iterations
        self.test_results = []
        self.window = None
        self.ServerChoosen = None
        self.InterationChoosen = None
        self.progress = None
        self.title("Bandwidth Test")
        self.geometry("800x600")
        self.create_widget()

    def create_widget(self):

        menubar = tk.Menu(self)

        option = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label='option', menu=option)
        option.add_command(label="BW Test", command=self.bandwidth_test)
        option.add_command(label="Export as Excel", command=self.export_bandwidth_test_to_excel)
        option.add_command(label="Configure Server", command=self.configure_setting)
        option.add_separator()
        option.add_command(label="Exit", command=self.destroy)

        self.config(menu=menubar)

    def run_iperf3_test(self, reverse):
        command = [
            'iperf3',
            '-c', self.server,
            '-p', str(self.port),
            '-t', str(self.duration),
            '-J',  # JSON output,
        ]

        if reverse:
            command.append('-R')

        try:
            result = subprocess.run(command, capture_output=True, text=True, check=True)
            result_data = json.loads(result.stdout)

            if 'error' in result_data:
                return self.test_results.append({'error': result_data['error']})
            else:
                for interval in result_data['intervals']:
                    if reverse:
                        self.test_results.append({
                            'received_Mbps': interval['sum']['bits_per_second'] / 1e6,
                        })
                    else:
                        self.test_results.append({
                            'sent_Mbps': interval['sum']['bits_per_second'] / 1e6,
                        })
        except subprocess.CalledProcessError as e:
            self.test_results.append({'error': str(e)})
        except json.JSONDecodeError:
            self.test_results.append({'error': 'Failed to parse JSON output from iperf3'})

    def run_multiple_tests(self, window):
        self.test_results.clear()
        self.upl.clear()
        self.dowl.clear()
        error_cnt = 0
        start_time = time.time()
        for i in range(self.iterations):
            self.progress['value'] = i * (100 / self.iterations)
            window.update_idletasks()
            self.run_iperf3_test(False)
            sleep(1)
            self.run_iperf3_test(True)

        end_time = time.time()
        print(end_time - start_time)
        self.progress['value'] = 100
        window.update_idletasks()
        for res in self.test_results:
            if 'sent_Mbps' in res:
                self.upl.append(round(res["sent_Mbps"], 1))
            elif 'received_Mbps' in res:
                self.dowl.append(round(res["received_Mbps"], 1))
            else:
                error_cnt += 1
        print(time.time() - end_time)
        if error_cnt >= (self.iterations / 5):
            window.message = messagebox.showerror(title="test state", message="test failed")
            window.destroy()
            return
        window.message = messagebox.showinfo(title="test state", message="successfully test")
        fig = Figure(figsize=(5, 5), dpi=80)
        plot1 = fig.add_subplot(111)

        plot1.plot(self.upl, label='Upload')
        plot1.plot(self.dowl, label='Download')
        plot1.legend()

        canvas = FigureCanvasTkAgg(fig, master=window)
        canvas.draw()
        canvas.get_tk_widget().pack()

        toolbar = NavigationToolbar2Tk(canvas, window)
        toolbar.update()
        canvas.get_tk_widget().pack()

        window.result_text = tk.Text(window, height=10, width=50)
        window.result_text.pack()
        average_upl, average_dowl = self.average_bandwidth()
        if average_upl == 'error' and average_dowl == ' error':
            window.message = messagebox.showerror(title="average bandwidth state", message="average bandwidth error")
            window.destroy()
            return
        server = {i for i in Server_List if Server_List[i] == f"{self.server}"}
        window.result_text.insert(tk.END, f"Upload: {average_upl} Mbps\n"
                                          f"Dowload: {average_dowl} Mbps\n"
                                          f"Server: {server}\n")

    def bandwidth_test(self):
        window = tk.Tk()
        window.title("Bandwidth Test Module")
        window.geometry('800x600')
        self.progress = ttk.Progressbar(window, orient="horizontal", length=100, mode='determinate')
        self.progress.pack(pady=10)

        start_button = tk.Button(window, text='Start', command=partial(self.run_multiple_tests, window))
        start_button.pack(pady=10)

    def export_bandwidth_test_to_excel(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Bandwidth Test Result"
        sheet['A1'].value = "Upload (Mbps)"
        sheet['B1'].value = "Download (Mbps)"
        row = 2
        for data in self.upl:
            sheet[f'A{row}'].value = data
            row += 1

        row = 2
        for data in self.dowl:
            sheet[f'B{row}'].value = data
            row += 1

        average_upl, average_dowl = self.average_bandwidth()
        sheet['D5'].value = "Average Upload: "
        sheet['D6'].value = "Average Download: "

        sheet['E5'].value = "{:.2f}".format(average_upl)
        sheet['E6'].value = "{:.2f}".format(average_dowl)

        # Create upload chart
        upload_chart = LineChart()
        upload_chart.title = "Upload Chart"
        upload_chart.x_axis.title = "Times"
        upload_chart.y_axis.title = "Mbps"
        upload_values = Reference(sheet, min_col=1, max_col=1, min_row=2, max_row=len(self.upl) + 1)
        upload_chart.add_data(upload_values, titles_from_data=True)
        sheet.add_chart(upload_chart, "G2")

        # Create download chart
        download_chart = LineChart()
        download_chart.title = "Download Chart"
        download_chart.x_axis.title = "Times"
        download_chart.y_axis.title = "Mbps"
        download_values = Reference(sheet, min_col=2, max_col=2, min_row=2, max_row=len(self.dowl) + 1)
        download_chart.add_data(download_values, titles_from_data=True)
        sheet.add_chart(download_chart, "G20")

        files = [('Excel Files', '*.xlsx')]
        save_path = asksaveasfilename(filetypes=files)
        if save_path:
            wb.save(save_path)
            messagebox.showinfo(title="Export state", message="Export Completed")

    def average_bandwidth(self):
        if len(self.upl) == 0 or len(self.dowl) == 0:
            return 'error', 'error'
        average_upl = sum(self.upl) / len(self.upl)
        average_dowl = sum(self.dowl) / len(self.dowl)
        return average_upl, average_dowl

    def configure_setting(self):
        window = tk.Tk()
        window.title("configure")
        window.geometry('640x480')

        tk.Label(window, text="Select the Server :",
                 font=("Times New Roman", 14)).grid(column=0,
                                                    row=15, padx=10, pady=25)
        var = tk.StringVar()
        var2 = tk.StringVar()
        self.ServerChoosen = ttk.Combobox(window, width=27, textvariable=var)
        self.ServerChoosen['values'] = ('Singapore', 'Tokyo', 'HongKong')

        self.ServerChoosen.grid(column=1, row=15)
        self.ServerChoosen.current(0)

        tk.Label(window, text="Select Interation :",
                 font=("Times New Roman", 14)).grid(column=0,
                                                    row=40, padx=10, pady=25)
        self.InterationChoosen = ttk.Combobox(window, width=27, textvariable=var2)
        self.InterationChoosen['values'] = ('1', '10', '20', '50', '100', '1000')

        self.InterationChoosen.grid(column=1, row=40)
        self.InterationChoosen.current(0)

        save_button = tk.Button(window, text="Save", command=partial(self.save_seletion, window))
        save_button.grid(column=1, row=45, padx=10, pady=10)

    def save_seletion(self, window):
        choosen_server = self.ServerChoosen.get()
        self.server = Server_List['{}'.format(choosen_server)]

        choosen_interation = self.InterationChoosen.get()
        self.iterations = Interation_List['{}'.format(choosen_interation)]
        window.destroy()


if __name__ == '__main__':
    app = BandwidthTest()
    app.mainloop()
