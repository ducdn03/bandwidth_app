import json
import subprocess
import time
from time import sleep
import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
from tkinter.filedialog import asksaveasfilename
import openpyxl
from matplotlib.backends.backend_tkagg import (FigureCanvasTkAgg, NavigationToolbar2Tk)
from matplotlib.figure import Figure
from openpyxl.chart import Reference, LineChart
import threading
from PIL import Image, ImageTk


CONST_ACCEPTED_RSSI = -50


class BandwidthTest(tk.Tk):
    def __init__(self, server='89.187.160.1', port=5201, duration=10, iterations=1, stream=10):
        super().__init__()
        self.upl = []
        self.dowl = []
        self.error_cnt = 0
        self.server = server
        self.port = port
        self.duration = duration
        self.iterations = iterations
        self.stream = stream
        self.test_results = []
        self.ServerChosen = None
        self.DurationChosen = None
        self.StreamChosen = None
        self.PortChosen = None
        self.main_frame = None
        self.title("Bandwidth Test")
        self.geometry("800x600")
        self.create_widget()

    def create_widget(self):
        menubar = tk.Menu(self)

        option = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label='Option', menu=option)
        option.add_command(label="New Window", command=self.new_window)
        option.add_command(label="BW Test", command=self.bandwidth_test)
        option.add_command(label="Export as Excel", command=self.export_bandwidth_test_to_excel)
        option.add_command(label="Configure Server", command=self.configure_setting)
        option.add_command(label="Power Wifi Test", command=self.testing_power_wifi)
        option.add_separator()
        option.add_command(label="Exit", command=self.destroy)

        self.config(menu=menubar)

        self.main_frame = ttk.Frame(self)
        self.main_frame.config()
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    @staticmethod
    def _get_frames(img):
        with Image.open(img) as gif:
            index = 0
            frames = []
            while True:
                try:
                    gif.seek(index)
                    frame = ImageTk.PhotoImage(gif)
                    frames.append(frame)
                except EOFError:
                    break

                index += 1
            return frames

    def _play_gif(self, label, frames):
        total_delay = 50
        delay_frame = 100
        for frame in frames:
            if not self._is_loading:
                return
            self.after(total_delay, self._next_frame, frame, label, frames)
            total_delay += delay_frame
        self.after(total_delay, self._next_frame, frame, label, frames, True)

    def _next_frame(self, frame, label, frames, restart=False):
        if restart:
            try:
                label.config()
            except tk.TclError:
                return
            self.after(1, self._play_gif, label, frames)
            return
        try:
            label.config(image=frame)
        except tk.TclError:
            return

    def loading(self):
        for child in self.main_frame.winfo_children():
            child.destroy()
        loading_label = tk.Label(
            self.main_frame,
            background='WHITE',
            border=0,
            highlightthickness=0
        )
        self._is_loading = True
        loading_label.pack()
        frames = self._get_frames('loading.gif')
        self._play_gif(loading_label, frames)
        return

    def stop_loading(self):
        self._is_loading = False

    @staticmethod
    def new_window():
        new = BandwidthTest()
        new.mainloop()

    def run_iperf3_test(self, reverse, stop_event):
        if stop_event.is_set():
            return

        command = [
            'iperf3',
            '-c', self.server,
            '-p', str(self.port),
            '-t', str(self.duration),
            '-P', str(self.stream),
            '-J',  # JSON output
        ]

        if reverse:
            command.append('-R')

        try:
            process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
            stdout, stderr = process.communicate()
            while not stop_event.is_set() and process.poll() is None:
                sleep(10)

            if stop_event.is_set():
                process.terminate()
                return

            result_data = json.loads(stdout)

            if 'error' in result_data:
                self.test_results.append({'error': result_data['error']})
                messagebox.showerror(title='test state', message=f"error: {result_data['error']}", parent=self)
            else:
                for interval in result_data['intervals']:
                    if reverse:
                        self.test_results.append({
                            'received_Mbps': interval['sum']['bits_per_second'] / 1e6,
                        })
                    else:
                        self.test_results.append({
                            'sent_Mbps': interval['sum']['bits_per_second'] / 1e6,
                        })
        except subprocess.CalledProcessError as e:
            self.test_results.append({'error': str(e)})
        except json.JSONDecodeError:
            self.test_results.append({'error': 'Failed to parse JSON output from iperf3'})

    def check_server_status(self, stop_event):
        command = [
            'ping',
            '-c', '1',  # Send only 1 packet
            self.server,
        ]

        while not stop_event.is_set():
            try:
                start_time = time.time()
                process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                stdout, stderr = process.communicate()
                while (time.time() - start_time) < 3:
                    if process.poll() is None:
                        sleep(1)

                if ((time.time() - start_time) >= 3) and process.poll() is None:
                    process.terminate()
                    stop_event.set()
                    return

                # Check if 'bytes from' is in the output to determine success
                if 'bytes from' in stdout:
                    sleep(10)
                else:
                    print("Server did not respond to ping")
                    stop_event.set()
                    return
            except subprocess.CalledProcessError as e:
                print(f"Server check error: {str(e)}")
                stop_event.set()
                return

    def run_multiple_tests(self):
        def test_wrapper():
            self.clear_test_results()

            stop_event = threading.Event()
            thread1 = threading.Thread(target=self.run_iperf3_test, args=(False, stop_event))
            thread2 = threading.Thread(target=self.run_iperf3_test, args=(True, stop_event))
            thread3 = threading.Thread(target=self.check_server_status, args=(stop_event,))

            thread1.start()
            thread3.start()

            thread1.join()
            sleep(1)
            thread2.start()
            thread2.join()

            stop_event.set()

            self.process_test_results()

            if self.is_test_bandwidth_fail():
                messagebox.showerror(title="Test State", message="Test failed", parent=self)
                self.display_graph_plot(upl=self.upl, dowl=self.dowl)
                self.stop_loading()
                return

            messagebox.showinfo(title="Test State", message="Test successfully completed", parent=self)
            self.stop_loading()
            self.display_graph_plot(upl=self.upl, dowl=self.dowl)

        threading.Thread(target=test_wrapper).start()
        self.loading()

    def export_bandwidth_test_to_excel(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Bandwidth Test Result"
        sheet['A1'].value = "Upload (Mbps)"
        sheet['B1'].value = "Download (Mbps)"
        row = 2
        for data in self.upl:
            sheet[f'A{row}'].value = data
            row += 1

        row = 2
        for data in self.dowl:
            sheet[f'B{row}'].value = data
            row += 1

        average_upl, average_dowl = self.average_bandwidth(upl=self.upl, dowl=self.dowl)
        sheet['D5'].value = "Average Upload: "
        sheet['D6'].value = "Average Download: "

        sheet['E5'].value = "{:.2f}".format(average_upl)
        sheet['E6'].value = "{:.2f}".format(average_dowl)

        # Create upload chart
        upload_chart = LineChart()
        upload_chart.title = "Upload Chart"
        upload_chart.x_axis.title = "Times"
        upload_chart.y_axis.title = "Mbps"
        upload_values = Reference(sheet, min_col=1, max_col=1, min_row=1, max_row=len(self.upl) + 1)
        upload_chart.add_data(upload_values, titles_from_data=True)
        sheet.add_chart(upload_chart, "G2")

        # Create download chart
        download_chart = LineChart()
        download_chart.title = "Download Chart"
        download_chart.x_axis.title = "Times"
        download_chart.y_axis.title = "Mbps"
        download_values = Reference(sheet, min_col=2, max_col=2, min_row=1, max_row=len(self.dowl) + 1)
        download_chart.add_data(download_values, titles_from_data=True)
        sheet.add_chart(download_chart, "G20")

        files = [('Excel Files', '*.xlsx')]
        save_path = asksaveasfilename(filetypes=files)
        if save_path:
            wb.save(save_path)
            messagebox.showinfo(title="Export state", message="Export Completed", parent=self)

    def display_graph_plot(self, upl, dowl):
        for widget in self.main_frame.winfo_children():
            widget.destroy()
        fig = Figure(figsize=(5, 5), dpi=80)
        plot1 = fig.add_subplot(111)

        plot1.plot(upl, label='Upload')
        plot1.plot(dowl, label='Download')
        plot1.legend()

        canvas = FigureCanvasTkAgg(fig, master=self.main_frame)
        canvas.draw()
        canvas.get_tk_widget().pack()
        plot1.set_xlabel('time (t)')
        plot1.set_ylabel('Speed (Mbps)')
        toolbar = NavigationToolbar2Tk(canvas, self.main_frame)
        toolbar.update()
        canvas.get_tk_widget().pack(side=tk.TOP, fill=tk.BOTH, expand=1)

        result_text = tk.Text(self.main_frame, height=10, width=50)
        result_text.pack(fill=tk.BOTH, expand=1)
        average_upl, average_dowl = self.average_bandwidth(upl=upl, dowl=dowl)
        if average_upl == 'error' and average_dowl == 'error':
            messagebox.showerror(title="Average Bandwidth State", message="Average bandwidth error", parent=self)
            return
        result_text.insert(tk.END, f"Upload: {average_upl} Mbps\n"
                                   f"Download: {average_dowl} Mbps\n"
                                   f"Server: {self.server}\n"
                                   f"Port: {self.port}\n"
                                   f"Stream: {self.stream}\n"
                                   f"Duration: {self.duration}\n")

    def bandwidth_test(self):
        self.clear_main_frame()

        start_button = tk.Button(self.main_frame, text='Start', command=self.run_multiple_tests)
        start_button.pack(pady=10)

    @staticmethod
    def average_bandwidth(upl, dowl):
        if len(upl) == 0 or len(dowl) == 0:
            return 'error', 'error'
        average_upl = sum(upl) / len(upl)
        average_dowl = sum(dowl) / len(dowl)
        return average_upl, average_dowl

    def testing_power_wifi(self):
        self.clear_main_frame()

        start_button = tk.Button(self.main_frame, text='Start', command=self.start_power_wifi_test)
        start_button.pack(pady=10)

    def start_power_wifi_test(self):
        def display_test_result():
            for result in self.test_pass:
                print(result)
                if '2.4Ghz rssi passed' in result:
                    rssi_2ghz_value = tk.Label(self.main_frame, text=result['2.4Ghz rssi passed'],
                                               font=("Times New Roman", 20), fg="green")
                    rssi_2ghz_value.grid(column=2, row=0)
                elif '5.0Ghz rssi passed' in result:
                    rssi_5ghz_value = tk.Label(self.main_frame, text=result['5.0Ghz rssi passed'],
                                               font=("Times New Roman", 20), fg="green")
                    rssi_5ghz_value.grid(column=2, row=1)
                elif '2.4Ghz bandwidth passed' in result:
                    bandwidth_2ghz_value = tk.Label(self.main_frame, text=result['2.4Ghz bandwidth passed'],
                                                    font=("Times New Roman", 20), fg="green")
                    bandwidth_2ghz_value.grid(column=2, row=2)
                elif '5.0Ghz bandwidth passed' in result:
                    bandwidth_5ghz_value = tk.Label(self.main_frame, text=result['5.0Ghz bandwidth passed'],
                                                    font=("Times New Roman", 20), fg="green")
                    bandwidth_5ghz_value.grid(column=2, row=3)

        def log_result():
            self.clear_main_frame()
            print("log")
            rssi_2ghz_label = tk.Label(self.main_frame, text="2.4 Ghz rssi test",
                                       font=("Times New Roman", 20))
            rssi_2ghz_label.grid(column=0, row=0)
            rssi_2ghz_result = tk.Label(self.main_frame, font=("Times New Roman", 20), text="FAILED", fg="red")
            rssi_2ghz_result.grid(column=1, row=0)

            rssi_5ghz_label = tk.Label(self.main_frame, text="5.0 Ghz rssi test",
                                       font=("Times New Roman", 20))
            rssi_5ghz_label.grid(column=0, row=1)
            rssi_5ghz_result = tk.Label(self.main_frame, font=("Times New Roman", 20), text="FAILED", fg="red")
            rssi_5ghz_result.grid(column=1, row=1)

            bandwidth_2ghz_label = tk.Label(self.main_frame, text="2.4 Ghz bandwidth test",
                                            font=("Times New Roman", 20))
            bandwidth_2ghz_label.grid(column=0, row=2)
            bandwidth_2ghz_result = tk.Label(self.main_frame, font=("Times New Roman", 20), text="FAILED", fg="red")
            bandwidth_2ghz_result.grid(column=1, row=2)

            bandwidth_5ghz_label = tk.Label(self.main_frame, text="5.0 Ghz bandwidth test",
                                            font=("Times New Roman", 20))
            bandwidth_5ghz_label.grid(column=0, row=3)
            bandwidth_5ghz_result = tk.Label(self.main_frame, font=("Times New Roman", 20), text="FAILED", fg="red")
            bandwidth_5ghz_result.grid(column=1, row=3)

            result_button = tk.Button(self.main_frame, text="See Result", command=display_test_result)
            result_button.grid(column=1, row=4)

            for result in self.test_pass:
                if '2.4Ghz rssi passed' in result:
                    rssi_2ghz_result.config(text="PASS", fg="green")
                elif '5.0Ghz rssi passed' in result:
                    rssi_5ghz_result.config(text="PASS", fg="green")
                elif '2.4Ghz bandwidth passed' in result:
                    if result['2.4Ghz bandwidth passed'] >= 75:
                        bandwidth_2ghz_result.config(text="PASS", fg="green")
                elif '5.0Ghz bandwidth passed' in result:
                    if result['5.0Ghz bandwidth passed'] >= 280:
                        bandwidth_5ghz_result.config(text="PASS", fg="green")

        def test_5ghz():
            rssi_5ghz_value = self.get_rssi_value()
            if rssi_5ghz_value >= CONST_ACCEPTED_RSSI:
                self.test_pass.append({'5.0Ghz rssi passed': rssi_5ghz_value})
            run_10minutes_bandwidth_test('5.0Ghz')

        def test_2ghz():
            rssi_value = self.get_rssi_value()
            if rssi_value >= CONST_ACCEPTED_RSSI:
                self.test_pass.append({'2.4Ghz rssi passed': rssi_value})
            run_10minutes_bandwidth_test('2.4Ghz')

        def run_10minutes_bandwidth_test(frequency):
            def test_wrapper():
                self.clear_test_results()
                print(frequency)
                stop_event = threading.Event()
                thread1 = threading.Thread(target=self.run_iperf3_test, args=(False, stop_event))
                thread2 = threading.Thread(target=self.run_iperf3_test, args=(True, stop_event))
                thread3 = threading.Thread(target=self.check_server_status, args=(stop_event,))

                thread1.start()
                thread3.start()

                thread1.join()
                sleep(1)
                thread2.start()
                thread2.join()

                stop_event.set()

                self.process_test_results()

                if self.is_test_bandwidth_fail():
                    messagebox.showerror("failed")
                    self.stop_loading()
                    return

                average_upl, average_dowl = self.average_bandwidth(self.upl, self.dowl)
                self.stop_loading()
                if frequency == '2.4Ghz':
                    self.test_pass.append({'2.4Ghz bandwidth passed': average_dowl})
                    messagebox.showinfo(message="test complete")
                    config_label = tk.Label(self.main_frame, text="wait for config the AP to 5.0Ghz ...")
                    config_label.pack()

                    start_button = tk.Button(self.main_frame, text="start", command=test_5ghz)
                    start_button.pack()
                elif frequency == '5.0Ghz':
                    self.test_pass.append({'5.0Ghz bandwidth passed': average_dowl})
                    messagebox.showinfo(message="test complete")
                    self.stop_loading()
                    log_result()
                return

            threading.Thread(target=test_wrapper).start()
            self.loading()

        self.clear_main_frame()

        self.test_pass = []

        test_2ghz()

    def configure_setting(self):
        self.clear_main_frame()

        tk.Label(self.main_frame, text="Enter Server IP:", font=("Times New Roman", 14)).grid(row=1, column=0)
        self.ServerChosen = tk.Entry(self.main_frame, font=("Times New Roman", 14))
        self.ServerChosen.grid(column=1, row=1)

        tk.Label(self.main_frame, text="Enter No. Duration:", font=("Times New Roman", 14)).grid(row=2, column=0)
        self.DurationChosen = tk.Entry(self.main_frame, font=("Times New Roman", 14))
        self.DurationChosen.grid(column=1, row=2)

        tk.Label(self.main_frame, text="Enter No. Stream:", font=("Times New Roman", 14)).grid(row=3, column=0)
        self.StreamChosen = tk.Entry(self.main_frame, font=("Times New Roman", 14))
        self.StreamChosen.grid(column=1, row=3)

        tk.Label(self.main_frame, text="Enter No. Port:", font=("Times New Roman", 14)).grid(row=4, column=0)
        self.PortChosen = tk.Entry(self.main_frame, font=("Times New Roman", 14))
        self.PortChosen.grid(column=1, row=4)

        save_button = tk.Button(self.main_frame, text="Save", command=self.save_selection)
        save_button.grid(column=1, row=5)

    @staticmethod
    def get_wifi_interface():
        command = [
            'lshw',
            '-C',
            'network',
            '-json'
        ]
        process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        stdout, stderr = process.communicate()
        result = json.loads(stdout)
        for res in result:
            if res["description"] == "Wireless interface":
                return res["logicalname"]

    def get_rssi_value(self):
        wireless_interface = self.get_wifi_interface()
        result = subprocess.run(['iwconfig', wireless_interface], capture_output=True, text=True)
        for line in result.stdout.splitlines():
            data = line.split()
            for word in data:
                if "Tx-Power" in word:
                    return int(word[9:])
        """
        wireless_interface = self.get_wifi_interface()
        result = subprocess.run(['iwconfig', wireless_interface], capture_output=True, text=True)
        line = next(line for line in result.stdout.splitlines() if 'Link' in line)
        line = line.replace('/100', '').replace('=', ' ')
        parts = line.split()
        signal = parts[5]
        return int(signal)
        """

    def clear_main_frame(self):
        for widget in self.main_frame.winfo_children():
            widget.destroy()

    def clear_test_results(self):
        self.test_results.clear()
        self.upl.clear()
        self.dowl.clear()
        self.error_cnt = 0

    def process_test_results(self):
        for result in self.test_results:
            if 'sent_Mbps' in result:
                self.upl.append(round(result['sent_Mbps'], 1))
            elif 'received_Mbps' in result:
                self.dowl.append(round(result['received_Mbps'], 1))
            elif 'server_status' in result:
                if result['server_status'] == 'down':
                    print("Server is down")
            elif 'error' in result:
                self.error_cnt += 1
                print('error', self.error_cnt)

    def is_test_bandwidth_fail(self):
        return (self.error_cnt >= self.duration/5) or len(self.upl) == 0 or len(self.dowl) == 0

    def save_selection(self):
        server = self.ServerChosen.get()
        duration = self.DurationChosen.get()
        stream = self.StreamChosen.get()
        port = self.PortChosen.get()

        if len(server) > 0:
            self.server = server
        if duration:
            self.duration = int(duration)
        if stream:
            self.stream = int(stream)
        if port:
            self.port = int(port)

        messagebox.showinfo(title='save state', message="save completed", parent=self)
        for widget in self.main_frame.winfo_children():
            widget.destroy()


if __name__ == '__main__':
    app = BandwidthTest()
    app.mainloop()
